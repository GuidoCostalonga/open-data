"""Estrae le serie degli archivi parziali (lotti 5, 9, 11, 15, 18) presenti nel repository.

Ogni foglio DATI_* ha le 16 colonne comuni (B:Q) e, da R, i campi originali della fonte.
Uso: python3 estrai_archivi_parziali.py <cartella con L5/ L9/ ...> <uscita.json>
"""
import glob, json, os, re, sys
import openpyxl

radice, uscita = sys.argv[1], sys.argv[2]
SALTA = re.compile(r'(code|CODE|Code \(|_ID$|^STRUCTURE|^ACTION$|^Year|^Note$|^Unit$|^Area$|Reporter Countries)', re.I)
out = {}
for f in sorted(glob.glob(radice + '/L*/PAESI/*/*.xlsx') + glob.glob(radice + '/L*/AGGREGATI/*/*.xlsx')):
    lotto = int(f.split('/')[len(radice.split('/'))].lstrip('L'))
    paese = f.split('/')[-2]
    wb = openpyxl.load_workbook(f, read_only=True)
    for s in wb.sheetnames:
        if not s.startswith('DATI_'):
            continue
        righe = wb[s].iter_rows(min_row=3, values_only=True)
        testa = next(righe)[1:]
        extra = [(i, h) for i, h in enumerate(testa) if i >= 16 and h and not SALTA.search(str(h))]
        for r in righe:
            r = r[1:]
            if not r or r[0] is None:
                continue
            ds, ds_n, anno, val, unita, stato, agg_f, url, ident = r[2], r[3], r[4], r[5], r[7], r[11], r[12], r[13], r[15]
            d = out.setdefault(f'{lotto}|{ds}', {'n': ds_n, 'url': url, 'agg': agg_f, 'serie': {}, 'terr': {}})
            d['terr'].setdefault(paese, r[1])
            se = d['serie'].get(ident)
            if se is None:
                etich = ' · '.join(str(r[i]) for i, h in extra if r[i] not in (None, '', 'Total', 'Not applicable', '_Z', '_T'))
                se = d['serie'][ident] = {'l': etich, 'u': unita, 'd': {}}
            se['d'].setdefault(paese, {})[str(anno)] = [val, stato]
    wb.close()
    print(f, flush=True)
json.dump(out, open(uscita, 'w'))
print('fine', len(out))
